from pathlib import Path

from openpyxl import load_workbook

from table_report.cli import EXIT_USER_ERROR, run

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"


def test_end_to_end_on_examples(tmp_path, capsys):
    output = tmp_path / "report.xlsx"
    code = run([str(EXAMPLES), "-o", str(output), "-g", "город", "-s", "сумма", "-s", "Кол-во", "--dedupe"])
    out = capsys.readouterr().out

    assert code == 0
    assert "Файлов: 2, строк: 10, удалено дублей: 1" in out
    assert "не похожи на числа" in out  # "по запросу" in the February file
    wb = load_workbook(output)
    summary = [[c.value for c in row] for row in wb["Итоги"].iter_rows()]
    assert summary[0] == ["Город", "Строк", "Сумма", "Кол-во"]
    assert summary[1] == ["Рим", 3, 2535.5, 33]  # biggest first
    assert summary[-1] == ["ИТОГО", 10, 6005.7, 140]
    assert wb["Данные"].max_row == 11  # header + 10 rows


def test_unknown_column_is_a_user_error(tmp_path, capsys):
    code = run([str(EXAMPLES), "-o", str(tmp_path / "r.xlsx"), "-g", "Страна"])
    assert code == EXIT_USER_ERROR
    assert "Column 'Страна' not found" in capsys.readouterr().err
    assert not (tmp_path / "r.xlsx").exists()


def test_missing_file_is_a_user_error(tmp_path, capsys):
    assert run([str(tmp_path / "nope.csv")]) == EXIT_USER_ERROR
    assert "File not found" in capsys.readouterr().err


def test_telegram_requires_credentials(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)  # no .env here
    monkeypatch.delenv("BOT_TOKEN", raising=False)
    monkeypatch.delenv("CHAT_ID", raising=False)
    code = run([str(EXAMPLES), "-o", str(tmp_path / "r.xlsx"), "--telegram"])
    assert code == EXIT_USER_ERROR
    assert "BOT_TOKEN" in capsys.readouterr().err
