import json

import httpx
import pytest
from openpyxl import load_workbook

from table_report.telegram import TelegramError, send_document
from table_report.transform import summarize
from table_report.writer import write_report


def test_report_has_both_sheets_and_keeps_formulas_as_text(tmp_path):
    rows = [
        {"city": "Rome", "note": '=HYPERLINK("http://evil","click")', "sum": 10},
        {"city": "Milan", "note": "ok", "sum": 5},
    ]
    summary = summarize(rows, ["city"], ["sum"])
    path = write_report(tmp_path / "out" / "report.xlsx", ["city", "note", "sum"], rows, summary)

    wb = load_workbook(path)
    assert wb.sheetnames == ["Итоги", "Данные"]
    summary_sheet, data = wb["Итоги"], wb["Данные"]
    assert [c.value for c in summary_sheet[1]] == ["city", "Строк", "sum"]
    assert summary_sheet[1][0].font.bold
    assert [c.value for c in summary_sheet[summary_sheet.max_row]] == ["ИТОГО", 2, 15]
    evil = data["B2"]
    assert evil.value.startswith("=HYPERLINK") and evil.data_type == "s"  # text, not a live formula
    assert data.freeze_panes == "A2"


def _client(handler) -> httpx.Client:
    return httpx.Client(transport=httpx.MockTransport(handler))


def test_send_document_posts_file(tmp_path):
    report = tmp_path / "r.xlsx"
    report.write_bytes(b"data")
    seen = {}

    def handler(request: httpx.Request) -> httpx.Response:
        seen["url"] = str(request.url)
        seen["body"] = request.content
        return httpx.Response(200, json={"ok": True, "result": {}})

    send_document("123:SECRET", "42", report, "caption", client=_client(handler))
    assert seen["url"].endswith("/bot123:SECRET/sendDocument")
    assert b'name="chat_id"' in seen["body"] and b"r.xlsx" in seen["body"]


def test_send_document_error_never_leaks_token(tmp_path):
    report = tmp_path / "r.xlsx"
    report.write_bytes(b"data")

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(400, content=json.dumps({"ok": False, "description": "Bad Request: chat not found"}))

    with pytest.raises(TelegramError) as info:
        send_document("123:SECRET", "42", report, client=_client(handler))
    assert "chat not found" in str(info.value)
    assert "SECRET" not in str(info.value)


def test_network_error_is_wrapped(tmp_path):
    report = tmp_path / "r.xlsx"
    report.write_bytes(b"data")

    def handler(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("boom", request=request)

    with pytest.raises(TelegramError, match="Network error"):
        send_document("123:SECRET", "42", report, client=_client(handler))
