"""Reading CSV and Excel files into one common format."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

Row = dict[str, Any]

# Excel on Windows often saves CSV in cp1251; utf-8-sig also strips the BOM.
ENCODINGS = ("utf-8-sig", "cp1251")
SUPPORTED = (".csv", ".xlsx", ".xlsm")


@dataclass
class Table:
    source: str
    columns: list[str]
    rows: list[Row] = field(default_factory=list)


def collect_files(paths: list[Path]) -> list[Path]:
    """Expands directories into the supported files they contain; keeps order, drops duplicates."""
    result: list[Path] = []
    for path in paths:
        if path.is_dir():
            found = sorted(p for p in path.iterdir() if p.suffix.lower() in SUPPORTED and not p.name.startswith("~$"))
            result.extend(found)
        elif path.exists():
            result.append(path)
        else:
            raise FileNotFoundError(f"File not found: {path}")
    unique: list[Path] = []
    seen: set[Path] = set()
    for path in result:
        if path.resolve() not in seen:
            seen.add(path.resolve())
            unique.append(path)
    return unique


def read_table(path: Path) -> Table:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    raise ValueError(f"{path.name}: unsupported format {suffix!r} (use .csv or .xlsx)")


def _decode(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"{path.name}: cannot detect the text encoding")


def _read_csv(path: Path) -> Table:
    text = _decode(path)
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [row for row in csv.reader(io.StringIO(text), dialect) if any(cell.strip() for cell in row)]
    return _build(path.name, rows)


def _read_xlsx(path: Path) -> Table:
    workbook = load_workbook(path, read_only=True, data_only=True)  # data_only: values, not formulas
    try:
        sheet = workbook.worksheets[0]
        rows = [list(row) for row in sheet.iter_rows(values_only=True) if any(c not in (None, "") for c in row)]
    finally:
        workbook.close()
    return _build(path.name, rows)


def normalize_header(value: Any, index: int) -> str:
    text = " ".join(str(value).split()) if value is not None else ""
    return text or f"column_{index + 1}"


def clean_value(value: Any) -> Any:
    if isinstance(value, str):
        value = " ".join(value.split())
        return value or None
    return value


def _build(source: str, rows: list[list[Any]]) -> Table:
    if not rows:
        return Table(source=source, columns=[])
    header: list[str] = []
    for index, value in enumerate(rows[0]):
        name = normalize_header(value, index)
        # Two columns with the same title would overwrite each other in a dict row.
        candidate, n = name, 2
        while candidate.casefold() in (h.casefold() for h in header):
            candidate, n = f"{name} ({n})", n + 1
        header.append(candidate)

    body = []
    for raw in rows[1:]:
        values = list(raw)[: len(header)] + [None] * (len(header) - len(raw))
        body.append({name: clean_value(value) for name, value in zip(header, values, strict=True)})
    return Table(source=source, columns=header, rows=body)
