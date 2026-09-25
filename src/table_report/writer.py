"""Writing a formatted Excel report."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from table_report.readers import Row
from table_report.transform import TOTAL_LABEL, Summary

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
TOTAL_FONT = Font(bold=True)
MAX_WIDTH = 50


def _set(cell, value: Any) -> None:
    cell.value = value
    # openpyxl stores any string starting with "=" as a formula. Data from input files must
    # stay plain text, otherwise a cell like "=HYPERLINK(...)" becomes live in the report.
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def _fill_sheet(sheet: Worksheet, header: list[str], rows: list[list[Any]]) -> None:
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for values in rows:
        sheet.append([None] * len(values))
        for cell, value in zip(sheet[sheet.max_row], values, strict=True):
            _set(cell, value)

    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.iter_cols(min_row=1, max_row=min(sheet.max_row, 500)), start=1):
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(index)].width = min(longest + 2, MAX_WIDTH)


def write_report(path: Path, columns: list[str], rows: list[Row], summary: Summary) -> Path:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Итоги"
    _fill_sheet(summary_sheet, summary.header, summary.rows)
    last = summary_sheet[summary_sheet.max_row]
    if summary.rows and last[0].value == TOTAL_LABEL:
        for cell in last:
            cell.font = TOTAL_FONT
        summary_sheet.auto_filter.ref = f"A1:{get_column_letter(len(summary.header))}{summary_sheet.max_row - 1}"

    data_sheet = workbook.create_sheet("Данные")
    _fill_sheet(data_sheet, columns, [[row.get(c) for c in columns] for row in rows])

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
